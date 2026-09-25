"""
generador_imagenes.py
=====================
Genera imagenes lado a lado:
  - IZQUIERDA : el codigo HTML con resaltado de sintaxis (estilo Dracula oscuro)
  - DERECHA   : captura de pantalla real del archivo en el navegador Chromium

Ademas, inserta automaticamente la imagen en el documento Word
docs/JCCP_Web_Practicas.docx en la seccion correspondiente a cada practica
(justo despues de "Capturas de pantalla:"), y guarda una copia JCCP_Web_Act6.docx.

Uso:
    py generador_imagenes.py                       -> procesa todos los archivos de act6/
    py generador_imagenes.py act6/practica14.html  -> solo ese archivo
    py generador_imagenes.py act5/                 -> todos los archivos de act5/

Las imagenes se guardan en capturas/<nombre_sin_extension>.png
"""

import sys
import re
import copy
import tempfile
from pathlib import Path
from io import BytesIO

import docx
from docx.shared import Cm
from PIL import Image, ImageDraw, ImageFont
from pygments import highlight
from pygments.lexers import HtmlLexer
from pygments.formatters import ImageFormatter
from playwright.sync_api import sync_playwright

# ------------------------------------
# CONFIGURACION VISUAL
# ------------------------------------
BROWSER_WIDTH  = 900
BROWSER_HEIGHT = 600
CODE_WIDTH     = 900
PADDING        = 30
GAP            = 20
FONT_SIZE      = 15

BG_COLOR      = (18, 18, 24)
ACCENT_COLOR  = (94, 129, 238)
HEADER_HEIGHT = 50

OUTPUT_DIR  = Path("capturas")
DOCX_MASTER = Path("docs/JCCP_Web_Practicas.docx")
DOCX_ACT6   = Path("docs/JCCP_Web_Act6.docx")

# Ancho de imagen en el documento Word (cm)
IMG_CM = 15.0


# ------------------------------------
# GENERACION DE IMAGEN
# ------------------------------------

def generar_imagen_codigo(html_code: str, width: int) -> Image.Image:
    formatter = ImageFormatter(
        style="dracula",
        font_name="Consolas",
        font_size=FONT_SIZE,
        line_numbers=True,
        line_number_bg="#1e1e2e",
        line_number_fg="#7f849c",
        image_pad=12,
    )
    png_bytes = highlight(html_code, HtmlLexer(), formatter)
    img = Image.open(BytesIO(png_bytes)).convert("RGBA")
    if img.width > width:
        ratio = width / img.width
        img = img.resize((width, int(img.height * ratio)), Image.LANCZOS)
    return img


def capturar_navegador(html_path: Path, width: int, height: int) -> Image.Image:
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True)
        page = browser.new_page(viewport={"width": width, "height": height})
        page.goto(html_path.resolve().as_uri())
        page.wait_for_timeout(600)
        screenshot_bytes = page.screenshot(full_page=False)
        browser.close()
    return Image.open(BytesIO(screenshot_bytes)).convert("RGBA")


def agregar_label(img: Image.Image, text: str, color: tuple) -> Image.Image:
    bar_h = 32
    out = Image.new("RGBA", (img.width, img.height + bar_h), (0, 0, 0, 0))
    bar = Image.new("RGBA", (img.width, bar_h), color + (255,))
    draw = ImageDraw.Draw(bar)
    try:
        font = ImageFont.truetype("arial.ttf", 14)
    except Exception:
        font = ImageFont.load_default()
    draw.text((10, 7), text, fill=(255, 255, 255), font=font)
    out.paste(bar, (0, 0))
    out.paste(img, (0, bar_h))
    return out


def componer_imagen(code_img: Image.Image, browser_img: Image.Image, title: str) -> Image.Image:
    max_h = max(code_img.height, browser_img.height)

    def pad_bottom(im, h):
        if im.height >= h:
            return im
        c = Image.new("RGBA", (im.width, h), (30, 30, 40, 255))
        c.paste(im, (0, 0))
        return c

    code_img    = pad_bottom(code_img,    max_h)
    browser_img = pad_bottom(browser_img, max_h)

    total_w = PADDING + code_img.width + GAP + browser_img.width + PADDING
    total_h = HEADER_HEIGHT + PADDING + max_h + PADDING

    canvas = Image.new("RGBA", (total_w, total_h), BG_COLOR + (255,))
    draw = ImageDraw.Draw(canvas)

    draw.rectangle([(0, 0), (total_w, HEADER_HEIGHT)], fill=ACCENT_COLOR + (255,))
    try:
        font_title = ImageFont.truetype("arialbd.ttf", 20)
    except Exception:
        try:
            font_title = ImageFont.truetype("arial.ttf", 20)
        except Exception:
            font_title = ImageFont.load_default()
    draw.text((PADDING, 13), title, fill=(255, 255, 255), font=font_title)

    x_sep = PADDING + code_img.width + GAP // 2
    draw.line([(x_sep, HEADER_HEIGHT + 5), (x_sep, total_h - 5)], fill=ACCENT_COLOR + (180,), width=2)

    canvas.paste(code_img,    (PADDING,                        HEADER_HEIGHT + PADDING), code_img)
    canvas.paste(browser_img, (PADDING + code_img.width + GAP, HEADER_HEIGHT + PADDING), browser_img)

    return canvas.convert("RGB")


# ------------------------------------
# INSERCION EN WORD
# ------------------------------------

def _numero_practica(stem: str) -> str:
    """Extrae el numero de practica del nombre del archivo (ej. 'practica11' -> '11')."""
    m = re.search(r'\d+', stem)
    return m.group() if m else ""


def insertar_imagen_en_docx(doc: docx.Document, num: str, img_path: Path):
    """
    Busca la seccion 'Practica <num>' en el documento, localiza el parrafo
    'Capturas de pantalla:' y reemplaza los parrafos vacios de esa seccion
    por la imagen. Si ya habia una imagen, la reemplaza.
    """
    body = doc._body._body
    elements = list(body)

    # Encontrar indice del parrafo de titulo de esta practica
    title_idx = None
    for i, child in enumerate(elements):
        tag = child.tag.split('}')[-1]
        if tag != 'p':
            continue
        p = docx.text.paragraph.Paragraph(child, doc)
        # Coincide si el texto del parrafo contiene "Practica <num>" o "Practica<num>"
        text = p.text
        if re.search(rf'[Pp]r[aá]ctica\s*{num}\b', text):
            title_idx = i
            break

    if title_idx is None:
        print(f"     AVISO: No se encontro 'Practica {num}' en el documento.")
        return

    # Buscar el parrafo "Capturas de pantalla:" despues del titulo
    capt_idx = None
    next_title_idx = None
    for i in range(title_idx + 1, len(elements)):
        child = elements[i]
        tag = child.tag.split('}')[-1]
        if tag == 'p':
            p = docx.text.paragraph.Paragraph(child, doc)
            if 'Capturas de pantalla' in p.text:
                capt_idx = i
            # Detectar siguiente practica para no sobrepasar
            if p.style.name == 'Heading 1' and i > title_idx + 1:
                next_title_idx = i
                break

    if capt_idx is None:
        print(f"     AVISO: No se encontro 'Capturas de pantalla:' para Practica {num}.")
        return

    # Eliminar parrafos vacios y antiguas imagenes entre capt_idx+1 y next_title_idx (o tabla)
    end_idx = next_title_idx if next_title_idx else len(elements)
    to_remove = []
    for i in range(capt_idx + 1, end_idx):
        child = elements[i]
        tag = child.tag.split('}')[-1]
        if tag == 'p':
            p = docx.text.paragraph.Paragraph(child, doc)
            if p.text.strip() == '' or 'drawing' in child.xml:
                to_remove.append(child)
        elif tag == 'tbl':
            break  # no tocar la tabla de observaciones

    for el in to_remove:
        body.remove(el)

    # Insertar parrafo con la imagen justo despues de "Capturas de pantalla:"
    # Obtener el elemento actualizado de capt_idx
    elements = list(body)
    capt_el = None
    counter = 0
    for child in body:
        tag = child.tag.split('}')[-1]
        if tag == 'p':
            p = docx.text.paragraph.Paragraph(child, doc)
            if 'Capturas de pantalla' in p.text and counter > title_idx - 1:
                capt_el = child
                break
            if re.search(rf'[Pp]r[aá]ctica\s*{num}\b', p.text):
                counter = 1

    if capt_el is None:
        # fallback: recorrer de nuevo
        for child in body:
            tag = child.tag.split('}')[-1]
            if tag == 'p':
                p = docx.text.paragraph.Paragraph(child, doc)
                if 'Capturas de pantalla' in p.text:
                    capt_el = child

    if capt_el is None:
        print(f"     AVISO: No se pudo localizar el parrafo de capturas para Practica {num}.")
        return

    # Crear un parrafo nuevo con la imagen y lo insertamos despues de capt_el
    new_p = doc.add_paragraph()
    run = new_p.add_run()
    run.add_picture(str(img_path), width=Cm(IMG_CM))

    # Mover el nuevo parrafo justo despues de capt_el en el XML
    body.remove(new_p._element)
    capt_el.addnext(new_p._element)

    print(f"     Imagen insertada en Word para Practica {num}")


# ------------------------------------
# PROCESO PRINCIPAL POR ARCHIVO
# ------------------------------------

def procesar_archivo(html_path: Path, doc: docx.Document):
    """Genera la imagen y la inserta en el documento Word."""
    print(f"  -> Procesando: {html_path}")

    html_code = html_path.read_text(encoding="utf-8", errors="replace")
    stem      = html_path.stem
    num       = _numero_practica(stem)
    title     = stem.replace("_", " ").replace("practica", "Practica ").title()

    # Generar imagen compuesta
    code_img    = generar_imagen_codigo(html_code, CODE_WIDTH)
    code_img    = agregar_label(code_img,    "Codigo HTML",          (40, 44, 78))
    browser_img = capturar_navegador(html_path, BROWSER_WIDTH, BROWSER_HEIGHT)
    browser_img = agregar_label(browser_img, "Vista en el navegador", (30, 80, 60))
    final       = componer_imagen(code_img, browser_img, title)

    # Guardar PNG
    OUTPUT_DIR.mkdir(exist_ok=True)
    out_path = OUTPUT_DIR / f"{stem}.png"
    final.save(out_path, "PNG")
    print(f"     PNG guardado -> {out_path}")

    # Insertar en Word si tenemos numero de practica
    if num:
        insertar_imagen_en_docx(doc, num, out_path)


def resolver_targets():
    if len(sys.argv) < 2:
        return sorted(Path("act6").glob("*.html"))
    targets = []
    for arg in sys.argv[1:]:
        p = Path(arg)
        if p.is_dir():
            targets.extend(sorted(p.glob("*.html")))
        elif p.is_file() and p.suffix == ".html":
            targets.append(p)
        else:
            print(f"Advertencia: No se reconoce: {arg}")
    return targets


# ------------------------------------
# MAIN
# ------------------------------------
if __name__ == "__main__":
    targets = resolver_targets()

    if not targets:
        print("No se encontraron archivos HTML. Comprueba la ruta.")
        sys.exit(1)

    print(f"\nGenerando {len(targets)} imagen(es) en '{OUTPUT_DIR}/'...\n")

    # Cargar documento Word una sola vez
    doc = docx.Document(str(DOCX_MASTER))

    errores = []
    for html_path in targets:
        try:
            procesar_archivo(html_path, doc)
        except Exception as e:
            errores.append((html_path, e))
            print(f"     ERROR en {html_path}: {e}")

    # Guardar documento actualizado
    doc.save(str(DOCX_MASTER))
    doc.save(str(DOCX_ACT6))
    print(f"\nDocumento Word actualizado: {DOCX_MASTER}")
    print(f"Copia para entrega guardada: {DOCX_ACT6}")

    print(f"\nCompletado. {len(targets) - len(errores)}/{len(targets)} imagenes generadas.")
    if errores:
        print("Errores encontrados:")
        for p, e in errores:
            print(f"  {p}: {e}")
